import io
import os
import time
from collections.abc import Callable
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.basin_files import rectangle_geojson
from tests.test_basin_pages import YEARS, _spy_loads, _spy_store_access
from tests.test_point_result_pages import precipitation_reads  # noqa: F401 (fixture)

FORMAT_NOTE = (
    "CSV files are for R and Python. To open the data in Excel, use the .xlsx workbook: with "
    "Spanish regional settings, Excel does not split the columns of a comma-separated file."
)
FOUR_CELLS = (-76.30, -12.40, -76.10, -12.20)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def store_dir(tmp_path: Path) -> Path:
    return tmp_path / "store"


@pytest.fixture
def client(make_client: Callable, reference_dir: Path, store_dir: Path) -> TestClient:
    return make_client(reference_dir, store_dir)


@pytest.fixture
def url(client: TestClient) -> str:
    files = {"basin_file": ("basin.geojson", rectangle_geojson(*FOUR_CELLS), "application/json")}
    response = client.post("/basin", data=YEARS, files=files, follow_redirects=False)
    assert response.status_code == 303
    return response.headers["location"]


def test_download_links(client: TestClient, url: str) -> None:
    html = client.get(f"{url}?page=2").text
    for path in (
        "export/daily.csv",
        "export/annual.csv",
        "export/climatology.csv",
        "export/cells.csv",
    ):
        assert f'href="{url}/{path}"' in html
    assert f'href="{url}/export.xlsx"' in html
    assert "page=2" not in html.split('aria-label="Downloads"', 1)[1].split("</nav>", 1)[0]
    assert FORMAT_NOTE in " ".join(html.split())


@pytest.mark.parametrize(
    ("table", "first_line"),
    [
        ("daily", "date,precipitation_mm"),
        ("annual", "year,total_mm,max_24h_mm,valid_days,complete"),
        ("climatology", "month,mean_mm,sd_mm,n_years"),
        ("cells", "cell_lon,cell_lat,used,valid_days"),
    ],
)
def test_csv_downloads_without_reads(
    client: TestClient,
    url: str,
    table: str,
    first_line: str,
    precipitation_reads: list[tuple[int, ...]],  # noqa: F811
) -> None:
    precipitation_reads.clear()
    response = client.get(f"{url}/export/{table}.csv")
    assert response.status_code == 200
    assert response.headers["content-type"] == "text/csv; charset=utf-8"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="pisco_basin_-12.3_-76.2_1981-1985_{table}.csv"'
    )
    assert response.text.split("\n", 1)[0] == first_line
    assert precipitation_reads == []


def test_daily_csv_values(client: TestClient, url: str) -> None:
    frame = pd.read_csv(io.BytesIO(client.get(f"{url}/export/daily.csv").content))
    assert len(frame) == 1826
    assert frame.loc[frame["date"] == "1981-01-03", "precipitation_mm"].item() == 3.0


def test_workbook_download(client: TestClient, url: str) -> None:
    response = client.get(f"{url}/export.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX
    assert response.headers["content-disposition"] == (
        'attachment; filename="pisco_basin_-12.3_-76.2_1981-1985.xlsx"'
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["daily", "annual", "climatology", "cells", "metadata"]


def test_unknown_table(client: TestClient, url: str) -> None:
    assert client.get(f"{url}/export/location.csv").status_code == 404


def test_expired_result(client: TestClient, url: str, store_dir: Path) -> None:
    stamp = time.time() - 25 * 3600
    os.utime(store_dir / f"{url.rsplit('/', 1)[1]}.json", (stamp, stamp))
    assert client.get(f"{url}/export.xlsx").status_code == 404
    assert client.get(f"{url}/export/daily.csv").status_code == 404


def test_invalid_id_never_touches_the_disk(
    client: TestClient, url: str, store_dir: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    calls = _spy_store_access(monkeypatch, store_dir)
    loads = _spy_loads(monkeypatch)
    for path in ("chart/daily.png", "export.xlsx", "export/cells.csv"):
        assert client.get(f"/basin/result/not-a-uuid/{path}").status_code == 404
    assert calls == [] and loads == []
