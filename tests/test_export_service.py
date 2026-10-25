import ast
import io
import re
from dataclasses import replace
from datetime import UTC, date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from app.config import Settings
from app.repositories.pisco_repository import CellRef, PiscoRepository
from app.services import export_service
from app.services.analysis_service import (
    BasinAnalysis,
    PointAnalysis,
    analyze_basin,
    analyze_point,
)
from app.services.basin_geometry import parse_basin_file
from app.services.basin_service import extract_basin
from app.services.export_service import (
    ExportTable,
    export_filename,
    metadata_rows,
    table_frame,
    to_csv,
    to_xlsx,
)
from app.services.point_service import PointResult
from tests.basin_files import rectangle_geojson

ISO_DATE = r"^\d{4}-\d{2}-\d{2}$"


def _reference_series(start: str, end: str, missing: tuple[str, ...] = ()) -> pd.Series:
    """Rainfall `(day index % 7) * 1.5` mm from 1981-01-01, stored as float32 like the file."""
    days = pd.date_range(start, end, freq="D", name="date")
    offset = (days[0] - pd.Timestamp("1981-01-01")).days
    values = ((np.arange(len(days)) + offset) % 7 * 1.5).astype(np.float32).astype(float)
    series = pd.Series(values, index=days, name="precipitation_mm")
    series[list(missing)] = np.nan
    return series


def _analysis(series: pd.Series, start_year: int, end_year: int) -> PointAnalysis:
    return analyze_point(
        PointResult(
            lon=-76.1390222,
            lat=-12.2935861,
            cell=CellRef(lat_index=2, lon_index=2, lon=-76.15, lat=-12.25),
            distance_km=4.991,
            resolution_deg=0.1,
            start_year=start_year,
            end_year=end_year,
            period_start=series.index[0].date(),
            period_end=series.index[-1].date(),
            total_days=len(series),
            missing_days=int(series.isna().sum()),
            series=series,
        )
    )


@pytest.fixture(scope="module")
def five_years() -> PointAnalysis:
    return _analysis(
        _reference_series("1981-01-01", "1985-12-31", ("1981-01-03", "1981-06-01")), 1981, 1985
    )


@pytest.fixture(scope="module")
def gappy() -> PointAnalysis:
    series = _reference_series("1981-01-01", "1982-12-31")
    series["1982"] = np.nan
    return _analysis(series, 1981, 1982)


def _lines(analysis: PointAnalysis, table: ExportTable) -> list[str]:
    return to_csv(table_frame(table, analysis)).decode("utf-8").split("\n")


# --- CSV tables -----------------------------------------------------------------------------


def test_daily_csv(five_years: PointAnalysis) -> None:
    lines = _lines(five_years, ExportTable.DAILY)
    assert lines[0] == "date,precipitation_mm"
    assert lines[-1] == ""  # trailing newline
    rows = lines[1:-1]
    assert len(rows) == 1826
    assert rows[:3] == ["1981-01-01,0.0", "1981-01-02,1.5", "1981-01-03,"]
    assert rows[-1] == "1985-12-31,7.5"


def test_annual_csv(five_years: PointAnalysis) -> None:
    lines = _lines(five_years, ExportTable.ANNUAL)
    assert lines[0] == "year,total_mm,max_24h_mm,valid_days,complete"
    assert len(lines[1:-1]) == 5
    assert "1981,1629.0,9.0,363,TRUE" in lines
    assert "1984,1648.5,9.0,366,TRUE" in lines


def test_annual_csv_year_without_data(gappy: PointAnalysis) -> None:
    assert "1982,,,0,FALSE" in _lines(gappy, ExportTable.ANNUAL)


def test_climatology_csv(five_years: PointAnalysis) -> None:
    lines = _lines(five_years, ExportTable.CLIMATOLOGY)
    assert lines[0] == "month,mean_mm,sd_mm,n_years"
    rows = lines[1:-1]
    assert [int(row.split(",")[0]) for row in rows] == list(range(1, 13))
    assert rows[0] == "1,137.7,6.6577,5"


def test_climatology_csv_single_year() -> None:
    analysis = _analysis(_reference_series("1985-01-01", "1985-12-31"), 1985, 1985)
    for row in _lines(analysis, ExportTable.CLIMATOLOGY)[1:-1]:
        _, _, sd_mm, n_years = row.split(",")
        assert (sd_mm, n_years) == ("", "1")


# --- CSV format -----------------------------------------------------------------------------


def test_single_precision_noise_removed() -> None:
    days = pd.date_range("1981-01-01", "1981-12-31", freq="D", name="date")
    series = pd.Series(np.full(len(days), 0.1, dtype=np.float32).astype(float), index=days)
    analysis = _analysis(series, 1981, 1981)
    assert _lines(analysis, ExportTable.DAILY)[1] == "1981-01-01,0.1"
    assert _lines(analysis, ExportTable.ANNUAL)[1] == "1981,36.5,0.1,365,TRUE"


@pytest.mark.parametrize("table", export_service.POINT_TABLES)
def test_no_encoding_artifacts(five_years: PointAnalysis, table: ExportTable) -> None:
    body = to_csv(table_frame(table, five_years))
    assert not body.startswith(b"\xef\xbb\xbf")
    assert b"\r" not in body
    assert b"nan" not in body.lower()
    assert b"None" not in body


def test_readable_with_default_parser(five_years: PointAnalysis) -> None:
    frames = {
        table: pd.read_csv(io.BytesIO(to_csv(table_frame(table, five_years))))
        for table in export_service.POINT_TABLES
    }
    daily, annual, climate = frames.values()
    assert list(daily.columns) == ["date", "precipitation_mm"]
    assert daily["date"].str.match(ISO_DATE).all()
    assert daily["precipitation_mm"].dtype == np.float64
    assert daily["precipitation_mm"].isna().sum() == 2
    assert {c: annual[c].dtype.kind for c in annual.columns} == {
        "year": "i",
        "total_mm": "f",
        "max_24h_mm": "f",
        "valid_days": "i",
        "complete": "b",
    }
    assert {c: climate[c].dtype.kind for c in climate.columns} == {
        "month": "i",
        "mean_mm": "f",
        "sd_mm": "f",
        "n_years": "i",
    }


def test_frames_share_the_rounding(five_years: PointAnalysis) -> None:
    frame = table_frame(ExportTable.CLIMATOLOGY, five_years)
    assert frame.loc[0, "sd_mm"] == 6.6577
    assert frame.loc[0, "month"] == 1
    assert isinstance(table_frame(ExportTable.DAILY, five_years).loc[0, "date"], date)


def test_module_imports_neither_fastapi_nor_xarray() -> None:
    tree = ast.parse(Path(export_service.__file__).read_text(encoding="utf-8"))
    modules = {
        (node.module or "").split(".")[0]
        for node in ast.walk(tree)
        if isinstance(node, ast.ImportFrom)
    } | {
        alias.name.split(".")[0]
        for node in ast.walk(tree)
        if isinstance(node, ast.Import)
        for alias in node.names
    }
    assert not modules & {"fastapi", "starlette", "xarray"}


# --- workbook -------------------------------------------------------------------------------

GENERATED_AT = datetime(2026, 10, 25, 12, tzinfo=UTC)
METADATA_FIELDS = [
    "source",
    "attribution",
    "requested_lon",
    "requested_lat",
    "cell_lon",
    "cell_lat",
    "distance_km",
    "start_year",
    "end_year",
    "period_start",
    "period_end",
    "completeness_rule",
    "generated_at_utc",
]


def _sheets(analysis: PointAnalysis) -> dict[str, pd.DataFrame]:
    return pd.read_excel(io.BytesIO(to_xlsx(analysis, GENERATED_AT)), sheet_name=None)


def _metadata(analysis: PointAnalysis) -> dict[str, object]:
    sheet = _sheets(analysis)["metadata"]
    assert list(sheet.columns) == ["field", "value"]
    return dict(zip(sheet["field"], sheet["value"], strict=True))


def test_workbook_sheets(five_years: PointAnalysis) -> None:
    sheets = _sheets(five_years)
    assert list(sheets) == ["daily", "annual", "climatology", "metadata"]
    for table in export_service.POINT_TABLES:
        assert list(sheets[table.value].columns) == export_service.COLUMNS[table]
    assert len(sheets["climatology"]) == 12
    assert sheets["climatology"].loc[0, "sd_mm"] == 6.6577


def test_workbook_daily_dates(five_years: PointAnalysis) -> None:
    daily = load_workbook(io.BytesIO(to_xlsx(five_years, GENERATED_AT)))["daily"]
    assert daily.max_row == 1827  # header + 1826 days
    assert daily["A2"].value.date() == date(1981, 1, 1)
    assert daily["A2"].number_format == "YYYY-MM-DD"
    assert daily["A4"].value.date() == date(1981, 1, 3)
    assert daily["B4"].value is None


def test_workbook_annual(five_years: PointAnalysis, gappy: PointAnalysis) -> None:
    annual = _sheets(five_years)["annual"]
    first = annual.iloc[0].tolist()
    assert first == [1981, 1629.0, 9.0, 363, True]
    assert annual["complete"].dtype == bool
    row_1982 = _sheets(gappy)["annual"].iloc[1]
    assert np.isnan(row_1982["total_mm"]) and np.isnan(row_1982["max_24h_mm"])
    assert (row_1982["valid_days"], row_1982["complete"]) == (0, False)


def test_workbook_metadata(five_years: PointAnalysis) -> None:
    rows = metadata_rows(five_years, GENERATED_AT)
    assert [field for field, _ in rows] == METADATA_FIELDS
    metadata = _metadata(five_years)
    assert "SENAMHI PISCOp v3.0" in str(metadata["source"])
    assert "SENAMHI" in str(metadata["attribution"])
    assert (metadata["cell_lon"], metadata["cell_lat"]) == (-76.15, -12.25)
    assert metadata["period_start"] == "1981-01-01"
    assert "95%" in str(metadata["completeness_rule"])
    assert metadata["generated_at_utc"] == "2026-10-25T12:00:00Z"
    assert re.match(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$", str(metadata["generated_at_utc"]))


def test_workbook_period_cut_by_end_of_file() -> None:
    analysis = _analysis(_reference_series("1985-01-01", "1985-06-30"), 1985, 1985)
    metadata = _metadata(analysis)
    assert (metadata["end_year"], metadata["period_end"]) == (1985, "1985-06-30")


# --- file names -----------------------------------------------------------------------------


def test_file_names(five_years: PointAnalysis) -> None:
    result = five_years.result
    assert export_filename(result, ExportTable.DAILY) == "pisco_-12.25_-76.15_1981-1985_daily.csv"
    assert export_filename(result, None) == "pisco_-12.25_-76.15_1981-1985.xlsx"


def test_file_name_strips_zeros_only() -> None:
    analysis = _analysis(_reference_series("1981-01-01", "1981-12-31"), 1981, 1981)
    result = replace(analysis.result, cell=CellRef(0, 0, lon=-70.0, lat=-0.05))
    assert export_filename(result, ExportTable.ANNUAL) == "pisco_-0.05_-70_1981-1981_annual.csv"


# --- basin exports ---------------------------------------------------------------------------


def _basin(
    directory: Path, rect: tuple[float, float, float, float], years: tuple[int, int] = (1981, 1985)
) -> BasinAnalysis:
    repo = PiscoRepository.open(Settings(data_dir=directory))
    try:
        basin = parse_basin_file("basin.geojson", rectangle_geojson(*rect))
        return analyze_basin(extract_basin(repo, basin, *years))
    finally:
        repo.close()


FOUR_CELLS = (-76.30, -12.40, -76.10, -12.20)
OCEAN_BASIN = (-76.40, -12.50, -76.20, -12.30)
SMALL_BASIN = (-76.14, -12.24, -76.12, -12.22)


def test_basin_daily_and_annual_csv(basin_dir: Path) -> None:
    analysis = _basin(basin_dir, FOUR_CELLS)
    daily = _lines(analysis, ExportTable.DAILY)
    assert daily[0] == "date,precipitation_mm"
    assert len(daily) == 1826 + 2  # header, rows, trailing newline
    assert "1981-01-03,3.0" in daily
    assert "1981,1638.0,9.0,365,TRUE" in _lines(analysis, ExportTable.ANNUAL)


def test_basin_day_without_data(varied_dir: Path) -> None:
    assert "1981-02-01," in _lines(_basin(varied_dir, FOUR_CELLS), ExportTable.DAILY)


def test_basin_cells_csv(basin_dir: Path) -> None:
    lines = _lines(_basin(basin_dir, OCEAN_BASIN), ExportTable.CELLS)
    assert lines[0] == "cell_lon,cell_lat,used,valid_days"
    assert lines[1] == "-76.35,-12.35,TRUE,1826"
    assert "-76.35,-12.45,FALSE,0" in lines
    assert len([line for line in lines[1:] if line]) == 4


def test_fallback_cells_csv(basin_dir: Path) -> None:
    lines = _lines(_basin(basin_dir, SMALL_BASIN), ExportTable.CELLS)
    assert lines[1:] == ["-76.15,-12.25,TRUE,1824", ""]


def test_basin_workbook(basin_dir: Path) -> None:
    analysis = _basin(basin_dir, OCEAN_BASIN)
    workbook = load_workbook(io.BytesIO(to_xlsx(analysis, GENERATED_AT)))
    assert workbook.sheetnames == ["daily", "annual", "climatology", "cells", "metadata"]
    cells = list(workbook["cells"].iter_rows(values_only=True))
    assert cells[0] == ("cell_lon", "cell_lat", "used", "valid_days")
    assert (-76.35, -12.45, False, 0) in cells
    assert all(isinstance(row[2], bool) for row in cells[1:])


def test_basin_metadata(basin_dir: Path) -> None:
    rows = dict(metadata_rows(_basin(basin_dir, OCEAN_BASIN), GENERATED_AT))
    assert list(rows) == [
        "source",
        "attribution",
        "basin_file",
        "features",
        "basin_area_km2",
        "centroid_lon",
        "centroid_lat",
        "cells_used",
        "cells_excluded",
        "cell_selection",
        "areal_mean",
        "start_year",
        "end_year",
        "period_start",
        "period_end",
        "completeness_rule",
        "generated_at_utc",
    ]
    assert rows["basin_file"] == "basin.geojson"
    assert (rows["features"], rows["basin_area_km2"]) == (1, 481.2)
    assert (rows["cells_used"], rows["cells_excluded"]) == (3, 1)
    assert "95%" in str(rows["completeness_rule"])
    assert rows["generated_at_utc"] == "2026-10-25T12:00:00Z"


def test_fallback_recorded_in_metadata(basin_dir: Path) -> None:
    rows = dict(metadata_rows(_basin(basin_dir, SMALL_BASIN), GENERATED_AT))
    assert rows["cell_selection"] == "nearest cell to the basin centroid (no cell center inside)"


def test_basin_file_names(basin_dir: Path) -> None:
    result = _basin(basin_dir, FOUR_CELLS).result
    assert (
        export_filename(result, ExportTable.CELLS) == "pisco_basin_-12.3_-76.2_1981-1985_cells.csv"
    )
    assert export_filename(result, None) == "pisco_basin_-12.3_-76.2_1981-1985.xlsx"


def test_cells_table_is_basin_only(five_years: PointAnalysis) -> None:
    with pytest.raises(ValueError, match="not available"):
        table_frame(ExportTable.CELLS, five_years)
